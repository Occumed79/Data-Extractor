import csv
import io
import json
import os
import re
import uuid
from contextlib import contextmanager
from dataclasses import dataclass, asdict
from datetime import datetime
from urllib.parse import urljoin, urlparse, parse_qs, urlencode, urlunparse

import requests
from bs4 import BeautifulSoup
from flask import Flask, flash, jsonify, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import create_engine, text
from sqlalchemy.engine import Engine

from queue_config import get_queue

try:
    from playwright.sync_api import sync_playwright
    PLAYWRIGHT_AVAILABLE = True
except Exception:
    PLAYWRIGHT_AVAILABLE = False

APP_DIR = os.path.dirname(os.path.abspath(__file__))
LOCAL_DB_PATH = os.path.join(APP_DIR, 'instance', 'portal.db')
MAX_PAGES_DEFAULT = 25
USER_AGENT = (
    'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) '
    'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36'
)


def normalize_database_url(url: str) -> str:
    if url.startswith('postgres://'):
        return 'postgresql+psycopg2://' + url[len('postgres://'):]
    if url.startswith('postgresql://'):
        return 'postgresql+psycopg2://' + url[len('postgresql://'):]
    return url


DATABASE_URL = normalize_database_url(
    os.getenv('DATABASE_URL', f'sqlite:///{LOCAL_DB_PATH}')
)

app = Flask(__name__)
app.secret_key = os.getenv('FLASK_SECRET_KEY', 'provider-portal-secret')

_ENGINE: Engine | None = None


def get_engine() -> Engine:
    global _ENGINE
    if _ENGINE is None:
        connect_args = {'check_same_thread': False} if DATABASE_URL.startswith('sqlite') else {}
        if DATABASE_URL.startswith('sqlite'):
            os.makedirs(os.path.join(APP_DIR, 'instance'), exist_ok=True)
        _ENGINE = create_engine(DATABASE_URL, future=True, pool_pre_ping=True, connect_args=connect_args)
    return _ENGINE


@contextmanager
def db_conn(begin: bool = False):
    engine = get_engine()
    with (engine.begin() if begin else engine.connect()) as conn:
        yield conn


try:
    get_queue().connection.ping()
    QUEUE_AVAILABLE = True
except Exception:
    QUEUE_AVAILABLE = False


@dataclass
class ProviderRecord:
    source_site: str = ''
    source_url: str = ''
    listing_page_url: str = ''
    name: str = ''
    entity_type: str = ''
    specialty: str = ''
    hospital: str = ''
    address: str = ''
    city: str = ''
    state: str = ''
    zip_code: str = ''
    price: str = ''
    rating: str = ''
    review_count: str = ''
    next_available: str = ''
    insurance: str = ''
    profile_url: str = ''
    booking_url: str = ''
    procedure_name: str = ''
    phone: str = ''
    notes: str = ''
    raw_json: dict | None = None


class BrowserFetcher:
    def __init__(self):
        self.play = None
        self.browser = None
        self.context = None
        self.page = None

    def __enter__(self):
        if not PLAYWRIGHT_AVAILABLE:
            return self
        self.play = sync_playwright().start()
        self.browser = self.play.chromium.launch(headless=True)
        self.context = self.browser.new_context(user_agent=USER_AGENT)
        self.page = self.context.new_page()
        return self

    def __exit__(self, exc_type, exc, tb):
        try:
            if self.context:
                self.context.close()
            if self.browser:
                self.browser.close()
            if self.play:
                self.play.stop()
        except Exception:
            pass

    def get_html(self, url: str, wait_ms: int = 1500) -> str:
        if PLAYWRIGHT_AVAILABLE and self.page:
            self.page.goto(url, wait_until='domcontentloaded', timeout=45000)
            self.page.wait_for_timeout(wait_ms)
            try:
                self.page.mouse.wheel(0, 2500)
                self.page.wait_for_timeout(500)
            except Exception:
                pass
            return self.page.content()
        resp = requests.get(url, headers={'User-Agent': USER_AGENT}, timeout=30)
        resp.raise_for_status()
        return resp.text


def normalize_space(text: str) -> str:
    return re.sub(r'\s+', ' ', text or '').strip()


def text_or_empty(node):
    return normalize_space(node.get_text(' ', strip=True)) if node else ''


def infer_site_type(url: str) -> str:
    host = urlparse(url).netloc.lower()
    if 'mdsave.com' in host:
        return 'mdsave'
    if 'zocdoc.com' in host:
        return 'zocdoc'
    return 'generic'


def find_next_links(base_url: str, soup: BeautifulSoup) -> list[str]:
    links = []
    seen = set()
    patterns = ('next', '>', '›', '→')
    for a in soup.find_all('a', href=True):
        label = text_or_empty(a).lower()
        href = urljoin(base_url, a['href'])
        aria = (a.get('aria-label') or '').lower()
        rel = ' '.join(a.get('rel', [])).lower()
        if any(p == label for p in patterns) or 'next' in label or 'next' in aria or 'next' in rel:
            if href not in seen:
                seen.add(href)
                links.append(href)
        elif label.isdigit():
            if href not in seen:
                seen.add(href)
                links.append(href)
    return links


def mdsave_parse_cards(base_url: str, html: str) -> tuple[list[ProviderRecord], list[str]]:
    soup = BeautifulSoup(html, 'html.parser')
    records: list[ProviderRecord] = []

    candidates = []
    for selector in ['article', 'section', 'div']:
        for node in soup.select(selector):
            text_blob = text_or_empty(node)
            if not text_blob:
                continue
            if '$' in text_blob and ('learn more' in text_blob.lower() or 'provider' in text_blob.lower() or 'hospital' in text_blob.lower()):
                candidates.append(node)

    seen = set()
    for node in candidates:
        text_blob = text_or_empty(node)
        if len(text_blob) < 20:
            continue
        name = ''
        profile_url = ''
        link = node.find('a', href=True)
        if link:
            profile_url = urljoin(base_url, link['href'])
            name = text_or_empty(link)
        if not name:
            for heading in node.find_all(['h2', 'h3', 'h4', 'strong']):
                maybe = text_or_empty(heading)
                if maybe and maybe.lower() not in {'learn more', 'book now'}:
                    name = maybe
                    break
        price_match = re.search(r'\$\s?[\d,]+(?:\.\d{2})?', text_blob)
        price = price_match.group(0).replace(' ', '') if price_match else ''
        address_match = re.search(r'\b\d{1,6}\s+[^\n,]+(?:,\s*[^\n,]+){1,3}', text_blob)
        address = address_match.group(0) if address_match else ''
        specialty = ''
        procedure_name = ''
        parts = [normalize_space(p) for p in text_blob.split('  ') if normalize_space(p)]
        if len(parts) > 1:
            specialty = parts[1] if parts[1] != name else ''
        proc_match = re.search(r'(?:Procedure|Service)\s*:?\s*([^$\n]+)', text_blob, re.I)
        if proc_match:
            procedure_name = normalize_space(proc_match.group(1))
        entity_type = 'hospital' if '/hospitals/' in base_url or 'hospital' in text_blob.lower() else 'provider'
        key = (name, address, price, profile_url)
        if name and key not in seen:
            seen.add(key)
            records.append(ProviderRecord(
                source_site='mdsave',
                source_url=base_url,
                listing_page_url=base_url,
                name=name,
                entity_type=entity_type,
                specialty=specialty,
                address=address,
                price=price,
                profile_url=profile_url,
                procedure_name=procedure_name,
                raw_json={'text': text_blob[:2500]}
            ))

    if not records:
        for a in soup.find_all('a', href=True):
            href = urljoin(base_url, a['href'])
            text_val = text_or_empty(a)
            if ('/hospitals/' in href or '/procedures/' in href) and text_val:
                records.append(ProviderRecord(
                    source_site='mdsave',
                    source_url=base_url,
                    listing_page_url=base_url,
                    name=text_val,
                    entity_type='directory_entry',
                    profile_url=href,
                    raw_json={'anchor_text': text_val}
                ))

    return records, find_next_links(base_url, soup)


def zocdoc_parse_cards(base_url: str, html: str) -> tuple[list[ProviderRecord], list[str]]:
    soup = BeautifulSoup(html, 'html.parser')
    records: list[ProviderRecord] = []
    seen = set()

    card_nodes = soup.find_all(['article', 'div', 'section'])
    for node in card_nodes:
        text_blob = text_or_empty(node)
        if not text_blob:
            continue
        lower = text_blob.lower()
        if not any(k in lower for k in ['review', 'book', 'available', 'insurance', 'specialist', 'dr.']):
            continue
        if len(text_blob) < 30:
            continue

        name = ''
        profile_url = ''
        for a in node.find_all('a', href=True):
            href = urljoin(base_url, a['href'])
            text_val = text_or_empty(a)
            if '/doctor/' in href or '/practice/' in href or '/dentist/' in href:
                profile_url = href
                if text_val and len(text_val) < 120:
                    name = text_val
                    break
        if not name:
            for h in node.find_all(['h2', 'h3', 'h4', 'strong']):
                maybe = text_or_empty(h)
                if maybe and len(maybe) < 120:
                    name = maybe
                    break
        rating_match = re.search(r'(\d(?:\.\d)?)\s*(?:stars?|/5)', lower)
        reviews_match = re.search(r'(\d[\d,]*)\s+reviews?', lower)
        specialty = ''
        address = ''
        next_available = ''
        insurance = ''
        booking_url = ''
        for a in node.find_all('a', href=True):
            txt = text_or_empty(a).lower()
            href = urljoin(base_url, a['href'])
            if 'book' in txt or 'appointment' in txt:
                booking_url = href
                break
        for line in [normalize_space(x) for x in text_blob.split('  ') if normalize_space(x)]:
            if 'insurance' in line.lower() and not insurance:
                insurance = line
            if 'next available' in line.lower() and not next_available:
                next_available = line
            if re.search(r'\b[a-zA-Z\s]+,\s*[A-Z]{2}\b', line) and not address:
                address = line
        pieces = [p for p in re.split(r'\n|\|', text_blob) if normalize_space(p)]
        if len(pieces) > 1 and not specialty:
            specialty = normalize_space(pieces[1])
        key = (name, address, profile_url)
        if name and key not in seen:
            seen.add(key)
            records.append(ProviderRecord(
                source_site='zocdoc',
                source_url=base_url,
                listing_page_url=base_url,
                name=name,
                entity_type='doctor',
                specialty=specialty,
                address=address,
                rating=rating_match.group(1) if rating_match else '',
                review_count=reviews_match.group(1) if reviews_match else '',
                next_available=next_available,
                insurance=insurance,
                profile_url=profile_url,
                booking_url=booking_url,
                raw_json={'text': text_blob[:2500]}
            ))

    return records, find_next_links(base_url, soup)


def parse_detail_page(site_type: str, url: str, html: str) -> dict:
    soup = BeautifulSoup(html, 'html.parser')
    text_blob = text_or_empty(soup)
    phone = ''
    phone_match = re.search(r'\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text_blob)
    if phone_match:
        phone = phone_match.group(0)
    notes = text_or_empty(soup.find(['main', 'article']))[:1000]
    return {'phone': phone, 'notes': notes}


def enumerate_pages(seed_url: str, next_links: list[str], max_pages: int, visited: set[str]) -> list[str]:
    queue = []
    for link in next_links:
        if link not in visited and len(visited) + len(queue) < max_pages:
            queue.append(link)
    if len(queue) < 2:
        parsed = urlparse(seed_url)
        qs = parse_qs(parsed.query)
        current = 1
        for key in ('page', 'p'):
            if key in qs:
                try:
                    current = int(qs[key][0])
                except Exception:
                    current = 1
                break
        for i in range(current + 1, min(current + 4, max_pages + 1)):
            qs2 = dict(qs)
            qs2['page'] = [str(i)]
            new_url = urlunparse(parsed._replace(query=urlencode(qs2, doseq=True)))
            if new_url not in visited and new_url not in queue:
                queue.append(new_url)
    return queue


def crawl_site(seed_url: str, site_type: str, detail_mode: bool, max_pages: int) -> list[ProviderRecord]:
    parse_func = mdsave_parse_cards if site_type == 'mdsave' else zocdoc_parse_cards
    results: list[ProviderRecord] = []
    visited_pages: set[str] = set()
    queued = [seed_url]

    with BrowserFetcher() as fetcher:
        while queued and len(visited_pages) < max_pages:
            current = queued.pop(0)
            if current in visited_pages:
                continue
            visited_pages.add(current)
            try:
                html = fetcher.get_html(current)
            except Exception:
                continue
            cards, next_links = parse_func(current, html)
            results.extend(cards)
            for candidate in enumerate_pages(current, next_links, max_pages, visited_pages):
                if candidate not in queued:
                    queued.append(candidate)

        if detail_mode:
            for rec in results:
                target = rec.profile_url or rec.booking_url
                if not target:
                    continue
                try:
                    html = fetcher.get_html(target, wait_ms=1000)
                    detail = parse_detail_page(site_type, target, html)
                    rec.phone = rec.phone or detail.get('phone', '')
                    rec.notes = rec.notes or detail.get('notes', '')
                except Exception:
                    continue

    deduped = []
    seen = set()
    for rec in results:
        key = (rec.source_site, rec.name, rec.address, rec.profile_url, rec.price)
        if rec.name and key not in seen:
            seen.add(key)
            deduped.append(rec)
    return deduped


def init_db():
    jobs_sql = """
    CREATE TABLE IF NOT EXISTS jobs (
        id TEXT PRIMARY KEY,
        created_at TEXT NOT NULL,
        updated_at TEXT NOT NULL,
        status TEXT NOT NULL,
        site_type TEXT NOT NULL,
        urls_json TEXT NOT NULL,
        detail_mode INTEGER NOT NULL DEFAULT 0,
        max_pages INTEGER NOT NULL DEFAULT 25,
        error_text TEXT,
        result_count INTEGER NOT NULL DEFAULT 0
    )
    """
    providers_sql = """
    CREATE TABLE IF NOT EXISTS providers (
        id TEXT PRIMARY KEY,
        created_at TEXT NOT NULL,
        job_id TEXT NOT NULL,
        source_site TEXT,
        source_url TEXT,
        listing_page_url TEXT,
        name TEXT,
        entity_type TEXT,
        specialty TEXT,
        hospital TEXT,
        address TEXT,
        city TEXT,
        state TEXT,
        zip_code TEXT,
        price TEXT,
        rating TEXT,
        review_count TEXT,
        next_available TEXT,
        insurance TEXT,
        profile_url TEXT,
        booking_url TEXT,
        procedure_name TEXT,
        phone TEXT,
        notes TEXT,
        raw_json TEXT
    )
    """
    with db_conn(begin=True) as conn:
        conn.execute(text(jobs_sql))
        conn.execute(text(providers_sql))


def save_job(job_id: str, site_type: str, urls: list[str], detail_mode: bool, max_pages: int):
    now = datetime.utcnow().isoformat()
    with db_conn(begin=True) as conn:
        conn.execute(text(
            'INSERT INTO jobs (id, created_at, updated_at, status, site_type, urls_json, detail_mode, max_pages) '
            'VALUES (:id, :created_at, :updated_at, :status, :site_type, :urls_json, :detail_mode, :max_pages)'
        ), {
            'id': job_id,
            'created_at': now,
            'updated_at': now,
            'status': 'queued',
            'site_type': site_type,
            'urls_json': json.dumps(urls),
            'detail_mode': int(detail_mode),
            'max_pages': max_pages,
        })


def update_job(job_id: str, **fields):
    if not fields:
        return
    fields['updated_at'] = datetime.utcnow().isoformat()
    assignments = ', '.join([f'{key} = :{key}' for key in fields.keys()])
    fields['job_id'] = job_id
    with db_conn(begin=True) as conn:
        conn.execute(text(f'UPDATE jobs SET {assignments} WHERE id = :job_id'), fields)


def insert_records(job_id: str, records: list[ProviderRecord]):
    with db_conn(begin=True) as conn:
        conn.execute(text('DELETE FROM providers WHERE job_id = :job_id'), {'job_id': job_id})
        now = datetime.utcnow().isoformat()
        for rec in records:
            payload = asdict(rec)
            raw_json = json.dumps(payload.pop('raw_json', {}) or {})
            data = {
                'id': uuid.uuid4().hex,
                'created_at': now,
                'job_id': job_id,
                **payload,
                'raw_json': raw_json,
            }
            conn.execute(text(
                'INSERT INTO providers ('
                'id, created_at, job_id, source_site, source_url, listing_page_url, name, entity_type, specialty, hospital, '
                'address, city, state, zip_code, price, rating, review_count, next_available, insurance, profile_url, '
                'booking_url, procedure_name, phone, notes, raw_json'
                ') VALUES ('
                ':id, :created_at, :job_id, :source_site, :source_url, :listing_page_url, :name, :entity_type, :specialty, :hospital, '
                ':address, :city, :state, :zip_code, :price, :rating, :review_count, :next_available, :insurance, :profile_url, '
                ':booking_url, :procedure_name, :phone, :notes, :raw_json'
                ')'
            ), data)
        conn.execute(text(
            'UPDATE jobs SET result_count = :result_count, updated_at = :updated_at WHERE id = :job_id'
        ), {
            'result_count': len(records),
            'updated_at': datetime.utcnow().isoformat(),
            'job_id': job_id,
        })


def fetch_job(job_id: str):
    with db_conn() as conn:
        return conn.execute(text('SELECT * FROM jobs WHERE id = :job_id'), {'job_id': job_id}).mappings().fetchone()


def fetch_providers(job_id: str):
    with db_conn() as conn:
        return conn.execute(text('SELECT * FROM providers WHERE job_id = :job_id ORDER BY created_at DESC, id DESC'), {'job_id': job_id}).mappings().fetchall()


init_db()


@app.route('/', methods=['GET'])
def index():
    with db_conn() as conn:
        jobs = conn.execute(text('SELECT * FROM jobs ORDER BY created_at DESC LIMIT 10')).mappings().fetchall()
    return render_template('index.html', jobs=jobs, playwright_available=PLAYWRIGHT_AVAILABLE, queue_available=QUEUE_AVAILABLE)


@app.route('/start', methods=['POST'])
def start_job():
    urls_blob = request.form.get('urls', '').strip()
    site_type = request.form.get('site_type', 'auto').strip() or 'auto'
    detail_mode = bool(request.form.get('detail_mode'))
    try:
        max_pages = max(1, min(150, int(request.form.get('max_pages', MAX_PAGES_DEFAULT))))
    except Exception:
        max_pages = MAX_PAGES_DEFAULT

    urls = [u.strip() for u in urls_blob.splitlines() if u.strip()]
    if not urls:
        flash('Please paste at least one URL.')
        return redirect(url_for('index'))

    job_id = uuid.uuid4().hex[:12]
    save_job(job_id, site_type, urls, detail_mode, max_pages)

    from tasks import run_scrape_job

    if QUEUE_AVAILABLE:
        queue = get_queue()
        queue.enqueue(run_scrape_job, job_id, job_timeout=60 * 60)
    else:
        try:
            run_scrape_job(job_id)
        except Exception as exc:
            update_job(job_id, status='failed', error_text=str(exc)[:4000])
            flash(f'Scrape failed: {str(exc)[:200]}')

    return redirect(url_for('job_detail', job_id=job_id))

@app.route('/history', methods=['GET'])
def history():
    with db_conn() as conn:
        jobs = conn.execute(text('SELECT * FROM jobs ORDER BY created_at DESC')).mappings().fetchall()
    return render_template('history.html', jobs=jobs)


@app.route('/job/<job_id>', methods=['GET'])
def job_detail(job_id):
    job = fetch_job(job_id)
    if not job:
        return 'Job not found', 404
    providers = fetch_providers(job_id)
    return render_template('job_detail.html', job=job, providers=providers)


@app.route('/api/job/<job_id>', methods=['GET'])
def api_job(job_id):
    job = fetch_job(job_id)
    if not job:
        return jsonify({'error': 'not found'}), 404
    return jsonify(dict(job))


EXPORT_FIELDS = [
    'id', 'source_site', 'source_url', 'listing_page_url', 'name', 'entity_type', 'specialty', 'hospital',
    'address', 'city', 'state', 'zip_code', 'price', 'rating', 'review_count', 'next_available',
    'insurance', 'profile_url', 'booking_url', 'procedure_name', 'phone', 'notes'
]


def selected_provider_rows(job_id: str):
    ids = request.values.getlist('selected_ids')
    with db_conn() as conn:
        if ids:
            placeholders = ', '.join([f':id_{i}' for i in range(len(ids))])
            params = {'job_id': job_id, **{f'id_{i}': value for i, value in enumerate(ids)}}
            query = text(f'SELECT * FROM providers WHERE job_id = :job_id AND id IN ({placeholders}) ORDER BY created_at DESC, id DESC')
            rows = conn.execute(query, params).mappings().fetchall()
        else:
            rows = conn.execute(text('SELECT * FROM providers WHERE job_id = :job_id ORDER BY created_at DESC, id DESC'), {'job_id': job_id}).mappings().fetchall()
    return rows


@app.route('/export/<job_id>.json', methods=['GET', 'POST'])
def export_json(job_id):
    rows = selected_provider_rows(job_id)
    data = [{field: row[field] for field in EXPORT_FIELDS} for row in rows]
    payload = io.BytesIO(json.dumps(data, indent=2).encode('utf-8'))
    return send_file(payload, mimetype='application/json', as_attachment=True, download_name=f'{job_id}.json')


@app.route('/export/<job_id>.csv', methods=['GET', 'POST'])
def export_csv(job_id):
    rows = selected_provider_rows(job_id)
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=EXPORT_FIELDS)
    writer.writeheader()
    for row in rows:
        writer.writerow({field: row[field] for field in EXPORT_FIELDS})
    payload = io.BytesIO(buffer.getvalue().encode('utf-8'))
    return send_file(payload, mimetype='text/csv', as_attachment=True, download_name=f'{job_id}.csv')


@app.route('/export/<job_id>.xlsx', methods=['GET', 'POST'])
def export_xlsx(job_id):
    rows = selected_provider_rows(job_id)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Providers'
    ws.append(EXPORT_FIELDS)
    for row in rows:
        ws.append([row[field] for field in EXPORT_FIELDS])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(
        bio,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'{job_id}.xlsx'
    )


@app.route('/export/<job_id>.pdf', methods=['GET', 'POST'])
def export_pdf(job_id):
    rows = selected_provider_rows(job_id)
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=letter, leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = [Paragraph('Provider Data Portal Report', styles['Title']), Spacer(1, 12)]
    elements.append(Paragraph(f'Job ID: {job_id}', styles['Normal']))
    elements.append(Paragraph(f'Record count: {len(rows)}', styles['Normal']))
    elements.append(Spacer(1, 12))
    table_data = [['Name', 'Type', 'Specialty', 'Address', 'Price/Rating']]
    for row in rows[:150]:
        price_rating = row['price'] or row['rating'] or ''
        table_data.append([
            row['name'] or '', row['entity_type'] or '', row['specialty'] or '',
            (row['address'] or '')[:50], price_rating
        ])
    table = Table(table_data, repeatRows=1, colWidths=[100, 60, 100, 180, 70])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#78C0A8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#5E412F')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.HexColor('#FCEBB6')])
    ]))
    elements.append(table)
    doc.build(elements)
    bio.seek(0)
    return send_file(bio, mimetype='application/pdf', as_attachment=True, download_name=f'{job_id}.pdf')


@app.route('/healthz', methods=['GET'])
def healthz():
    return jsonify({'status': 'ok'})


if __name__ == '__main__':
    init_db()
    app.run(debug=False, host='0.0.0.0', port=int(os.getenv('PORT', '5000')))
